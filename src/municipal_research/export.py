from __future__ import annotations

import json
import math
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .storage import read_json

DATE_FIELDS = ["publication_date", "decision_date", "implementation_date", "in_force_by_date", "effective_date"]


def cell_value(value):
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", str(value))
    if len(value) > 32000:
        value = value[:31900] + " [DISPLAY TRUNCATED; complete value in results.json/audit.jsonl]"
    return value


def export_workbook(run: Path, output: Path | None = None) -> Path:
    data = read_json(run / "results.json")
    manifest = read_json(run / "manifest.json")
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, rows, headers):
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.append(headers)
        for row in rows:
            ws.append([cell_value(row.get(h)) for h in headers])
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in ws[1]:
            cell.font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="234E63")
        for index, header in enumerate(headers, 1):
            width = 22
            if header in {"rationale", "quote", "matched_text", "detail", "definition", "explanation", "reason"}:
                width = 65
            elif "url" in header or "path" in header or header in {"issues", "gaps", "value", "queries", "hits", "meeting_archives"}:
                width = 48
            elif "id" in header or header in {"chunk_id", "stage"}:
                width = 30
            elif header.startswith(("char_", "line_", "page_")):
                width = 15
            ws.column_dimensions[get_column_letter(index)].width = max(width, len(header) + 2)
        ws.row_dimensions[1].height = 32
        for row in ws.iter_rows(min_row=2):
            lines = max(
                sum(max(1, math.ceil(len(part) / max(10, ws.column_dimensions[cell.column_letter].width - 3))) for part in str(cell.value or "").split("\n"))
                for cell in row
            )
            ws.row_dimensions[row[0].row].height = min(300, max(32, lines * 15))
        ws.freeze_panes = "C2" if len(headers) > 2 else "A2"
        ws.auto_filter.ref = ws.dimensions
        if rows:
            table = Table(displayName=re.sub(r"[^A-Za-z0-9]", "", name) + "Table", ref=ws.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.print_title_rows = "1:1"

    sheet("Summary", data["summary"], [
        "municipality_id", "municipality", "category", "needs_review", "documents",
        "chunks_scanned", "chunks_classified", "verified_quotes", "coverage", "gaps",
        "meeting_archives", "source_urls", *[f + "s" for f in DATE_FIELDS], "scopes",
        *[f"field:{key}" for key in data["config"]["research"]["fields"]],
    ])
    classifications, evidence, attributes, passes = [], [], [], []
    source_map = {d["id"]: d for d in data["documents"]}
    for record in data["classifications"]:
        base = {"municipality_id": record["municipality_id"], "document_id": record["document_id"], "chunk_id": record["chunk_id"], "source_url": record["url"]}
        decision = record["decision"]
        classifications.append({**base, **decision, "method": record["method"], "needs_review": record["needs_review"], "agreement": record["agreement"]})
        for attr in decision["attributes"]:
            attributes.append({**base, **attr})
        stages = [("final", record)] + [(p["stage"], p) for p in record["passes"] if "quotes" in p]
        for stage, item in stages:
            for quote in item.get("quotes", []):
                for occurrence, location in enumerate(quote["locations"] or [{}], 1):
                    evidence.append({
                        **base, "stage": stage, "quote_index": quote["quote_index"], "purpose": quote["purpose"],
                        "quote": quote["quote"], "verified": quote["verified"], "occurrence": occurrence,
                        "match": location.get("match"), "char_start": location.get("start"), "char_end": location.get("end"),
                        **{key: location.get(key) for key in ["line_start", "line_end", "page_start", "page_end", "page_char_start", "page_char_end", "matched_text"]},
                        "text_path": source_map.get(record["document_id"], {}).get("text_path"), "issue": quote["issue"],
                    })
        for entry in record["passes"]:
            passes.append({**base, "stage": entry["stage"], **entry.get("decision", {}), **entry.get("verdict", {}), "issues": entry.get("issues"), "error": entry.get("error")})

    sheet("Classifications", classifications, ["municipality_id", "category", *DATE_FIELDS, "temporal_relation", "scope", "rationale", "method", "needs_review", "agreement", "source_url", "document_id", "chunk_id"])
    sheet("Evidence", evidence, ["municipality_id", "stage", "quote_index", "purpose", "verified", "quote", "match", "occurrence", "char_start", "char_end", "line_start", "line_end", "page_start", "page_end", "page_char_start", "page_char_end", "matched_text", "source_url", "text_path", "document_id", "chunk_id", "issue"])
    sheet("Attributes", attributes, ["municipality_id", "name", "value", "quote_indices", "source_url", "chunk_id"])
    sheet("Sources", data["documents"], ["municipality_id", "title", "url", "retrieved_at", "media_type", "extractor", "warnings", "raw_path", "text_path", "numbered_path", "raw_sha256", "text_sha256", "id"])
    sheet("Discovery", data["discovery"], ["municipality_id", "method", "outcome", "url", "final_url", "title", "parent", "score", "depth", "error"])
    sheet("Triage", data.get("triage", []), ["municipality_id", "status", "reason", "hits", "url", "document_id", "chunk_id"])
    sheet("Pending_Searches", data.get("pending_searches", []), ["municipality_id", "municipality", "status", "reason", "meeting_archives", "queries"])
    sheet("Pass_Audit", passes, ["municipality_id", "stage", "category", *DATE_FIELDS, "scope", "rationale", "supported", "explanation", "issues", "error", "source_url", "chunk_id"])
    sheet("Errors", data["errors"], ["municipality_id", "stage", "url", "detail"])
    audit_path = run / "audit.jsonl"
    audit = [json.loads(line) for line in audit_path.read_text(encoding="utf-8").splitlines()] if audit_path.exists() else []
    sheet("API_Usage", [e for e in audit if e.get("stage") == "llm"], ["timestamp", "role", "model", "cache_hit", "response_id", "usage", "status", "key"])
    shard_manifests = manifest.get("shard_manifests", [])
    if shard_manifests:
        sheet("Shard_Manifests", shard_manifests, ["shard_index", "status", "municipality_count", "manifest_path", "manifest_sha256"])
    methodology = [
        {"item": "Question", "definition": data["config"]["research"]["question"]},
        {"item": "Exclusive cutoff", "definition": data["config"]["research"]["cutoff"]},
        {"item": "Run kind", "definition": manifest["mode"]},
        {"item": "Chronology", "definition": "Publication, formal decision, implementation and in-force-by dates are recorded separately and are never substituted for one another."},
        {"item": "Language thresholds", "definition": "Svenska 1, Svenska som andraspråk 1, SFI, GERS B1, GERS B2 and qualitative requirements remain distinct; no equivalence is inferred."},
        {"item": "Location convention", "definition": "Zero-based Unicode characters, end-exclusive; one-based lines and physical PDF pages. HTML has no page number. Offsets refer to saved text.txt, not raw HTML byte positions."},
        {"item": "Quote matching", "definition": "Exact or whitespace-only; all occurrences within the supplied chunk are listed. No fuzzy matching. Quote presence is distinct from semantic support."},
        {"item": "Triage", "definition": "Deterministic high-recall local triage precedes API classification. Only relevant/uncertain chunks are sent to the API. Excluded chunks are retained in Triage and never imply absence."},
        {"item": "Summary rule", "definition": "One conclusive category is retained. Multiple conclusive categories or conflicting date/scope/attribute values produce Review. Missing evidence and search failures never imply No."},
        {"item": "Coverage", "definition": "Coverage is bounded and gap-aware, not automatically exhaustive. Failed/limited searches remain Pending_Searches; sparse PDFs require manual OCR/review."},
        {"item": "Audit files", "definition": "results.json, manifest.json, audit.jsonl, raw sources, cleaned text, chunks, triage, units and cached LLM request/response material support audit and reproduction."},
    ]
    methodology += [{"item": label, "definition": rule["definition"]} for label, rule in data["config"]["research"]["labels"].items()]
    sheet("Methodology", methodology, ["item", "definition"])
    output = output or run / "results.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    saved = load_workbook(output, read_only=True, data_only=False)
    if saved.sheetnames != wb.sheetnames:
        raise RuntimeError("Workbook round-trip changed the sheet inventory")
    for ws in saved:
        if any(cell.data_type in {"f", "e"} for row in ws for cell in row):
            raise RuntimeError("Unexpected Excel formula/error from untrusted text")
    saved.close()
    return output
