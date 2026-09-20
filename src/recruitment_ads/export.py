"""Portable research exports. All estimates are recruitment wording, never policy status."""

import json
from pathlib import Path

import pandas as pd
import pyarrow.parquet as pq

from .aggregate import CATEGORY_MEASURES, build_panel, measures, national_trends
from .deduplicate import assign_spells
from .download import sha256, write_json
from .employer import EmployerMaster
from .pipeline import read_config
from .schemas import SCHEMA
from .validate import af_comparison, false_negative_metrics, metrics, preserve_review

METHODOLOGY = [
    (
        "Construct",
        "Explicit Swedish-language requirements in recruitment advertisements. These outputs never establish formal municipal policy.",
    ),
    (
        "Primary population",
        "Direct municipality legal-entity organisation-number matches; one configured SSYK occupation group; elderly_care_context=yes; nonempty original description and dated ad ID.",
    ),
    (
        "Scope",
        "Pilot: ten selected municipalities, 2017/2021/2024. Pilot aggregates are not national estimates. Other years remain NOT_PROCESSED.",
    ),
    (
        "2017 gap",
        "The inspected 2017 enriched archive has no employer organisation numbers. Primary results are EMPLOYER_ID_GAP, not zero. Exact validated-name matches appear only in a labelled sensitivity table.",
    ),
    (
        "Employer master",
        "Current official municipal legal entities matched to SCB municipality identifiers; unknown historical validity endpoints stay blank. Applying this crosswalk historically assumes continuity, recorded per ad.",
    ),
    (
        "Taxonomy",
        "Each formal level remains separate. No SFI/course/CEFR equivalence is assumed. Category shares may overlap.",
    ),
    (
        "Requirement meaning",
        "Only required semantic hits enter the numerator. Preferred, training, descriptive, alternative-language, application-instruction, uncertain and irrelevant statements do not.",
    ),
    (
        "Uncertainty",
        "No detected requirement is not proof of absence. All rates are unvalidated classifier outputs until human review. Inspect uncertain hits and complete text in the negative sample.",
    ),
    (
        "Denominator",
        "Eligible ads, including ads with no language hit. Duplicate ad IDs are flagged and excluded after the first record; original rows are retained. Missing text is excluded and counted.",
    ),
    (
        "Vacancies",
        "Known positive integer counts only for primary vacancy weighting. Missing/zero/invalid counts are not one. Missing-as-one sensitivity is a separate column; inspect missing-count frequency.",
    ),
    (
        "Spells",
        "Possible repost clusters: same legal employer, municipality, SSYK and normalized title, >=90% token-trigram Jaccard similarity within 45 days of the first ad. Earliest eligible ad represents a spell; within-spell changes are flagged. 30-day sensitivity is separate.",
    ),
    (
        "Fixed weights",
        "Equal 0.5 weights for SSYK 5321 and 5330 by default, fixed across time. Configurable; standardised rates are missing if either occupation lacks a denominator.",
    ),
    (
        "Coverage",
        "GOOD >=20; MODERATE 5–19; SPARSE 1–4; NO_ADS 0 observed eligible ads. No-ad shares remain missing. NOT_PROCESSED and EMPLOYER_ID_GAP also remain missing.",
    ),
    (
        "Time",
        "2016–2025 are complete calendar years by design, not a guarantee of exhaustive Platsbanken coverage. 2026 and development samples are excluded from default complete-year trends. 2006–2015 is disabled pending separate comparability validation.",
    ),
    (
        "Validation",
        "Approximately 600 stratified ads plus a separate random 300 predicted-negative ads, subject to available strata. Empty human labels mean pending. Design weights are provided; partial/nonrandom review can bias metrics. Threshold 95% precision is a target, not a guarantee.",
    ),
    (
        "AF comparison",
        "Original-text classification is primary. Official must_have/nice_to_have languages are preserved. Missing fields are unknown, not false. Disagreement does not establish which classifier is correct.",
    ),
    (
        "Manual coding",
        "Edit validation.csv and false_negatives.csv, or import reviewed validation.xlsx using the validation command. Existing coding is not overwritten during resampling. Workbook re-export fails when a validation workbook already exists.",
    ),
    (
        "Linkage",
        "Future formal-policy data must be merged separately on municipality_id/year. This pipeline creates no formal-policy variables or inferred adoption dates.",
    ),
]


def prepare_tables(output, master_path, config_path):
    output = Path(output)
    config = read_config(config_path)
    records = pq.read_table(output / "ads_classified.parquet").to_pylist()
    spells = pq.read_table(output / "recruitment_spells.parquet").to_pylist()
    run = json.loads((output / "run.json").read_text(encoding="utf-8"))
    inventory = json.loads((output / "source_inventory.json").read_text(encoding="utf-8"))
    sources = {}
    for item in inventory:
        n = item["counts"]["source_records"]
        sources[item["year"]] = {
            **item["source"],
            "employer_identifier_gap": item["field_missing"].get("employer.organization_number", 0)
            == n,
        }
    master = EmployerMaster.load(master_path)
    municipalities = [r for r in master.rows if r["municipality_id"] in run["scope"]]
    scope = (
        "all_290_municipalities"
        if len(municipalities) == 290
        else f"pilot_{len(municipalities)}_municipalities"
    )
    panel = build_panel(records, municipalities, config["years"], sources, config["occupations"])
    panel.to_parquet(output / "municipality_year.parquet", index=False)
    spell_panel = build_panel(
        spells, municipalities, config["years"], sources, config["occupations"]
    )
    spell_panel["analysis_unit"] = "earliest_eligible_ad_per_spell"
    spell_panel.to_parquet(output / "municipality_year_spells.parquet", index=False)
    national, occupations = national_trends(
        records, sources, config["fixed_occupation_weights"], scope
    )
    national_spells, _ = national_trends(spells, sources, config["fixed_occupation_weights"], scope)
    validation = pd.read_csv(
        output / "validation.csv",
        dtype={"municipality_id": str, "ad_id": str},
        keep_default_na=False,
    )
    negatives = pd.read_csv(
        output / "false_negatives.csv",
        dtype={"municipality_id": str, "ad_id": str},
        keep_default_na=False,
    )
    af, discordant = af_comparison(records)
    employer = (
        pd.DataFrame(records, columns=SCHEMA.names)
        .groupby(["year", "employer_match_method", "employer_validity_basis"], dropna=False)
        .size()
        .reset_index(name="n_retained_ads")
    )
    audit = pd.DataFrame(
        [
            {
                "year": item["year"],
                "source_url": item["source"]["source_url"],
                "source_hash": item["source"]["source_hash"],
                "retrieved_at": item["source"]["retrieved_at"],
                "data_version": item["source"]["data_version"],
                **item["counts"],
                **{"missing_" + k: v for k, v in item["field_missing"].items()},
            }
            for item in inventory
        ]
    )
    categories = national[
        [
            c
            for c in ["year", "geographic_scope", *["share_" + c for c in CATEGORY_MEASURES]]
            if c in national
        ]
    ]
    sensitivities = []
    for label, candidates in (
        (
            "validated_names_plus_exact_context_yes",
            [
                {
                    **r,
                    "primary_eligible": r["sensitivity_eligible"]
                    and r["elderly_care_context"] == "yes"
                    and not r["duplicate_ad_id"],
                }
                for r in records
            ],
        ),
        (
            "exact_orgnr_including_mixed_uncertain_context",
            [
                {
                    **r,
                    "primary_eligible": r["sensitivity_eligible"]
                    and r["employer_match_method"] == "orgnr_exact"
                    and not r["duplicate_ad_id"],
                }
                for r in records
            ],
        ),
    ):
        for year in sorted(sources):
            sensitivities.append(
                {
                    "analysis": label,
                    "year": year,
                    "geographic_scope": scope,
                    **measures(
                        [r for r in candidates if r["primary_eligible"] and r["year"] == year]
                    ),
                }
            )
    _, spells30 = assign_spells(records, 30, config["text_similarity_threshold"])
    for year in sorted(sources):
        if not sources[year]["employer_identifier_gap"]:
            sensitivities.append(
                {
                    "analysis": "30_day_spells_earliest_ad",
                    "year": year,
                    "geographic_scope": scope,
                    **measures([r for r in spells30 if r["year"] == year]),
                }
            )
    methodology = [
        (
            k,
            v.replace(
                "ten selected municipalities", str(len(municipalities)) + " selected municipalities"
            ),
        )
        for k, v in METHODOLOGY
    ]
    gaps = [str(y) for y, source in sources.items() if source["employer_identifier_gap"]]
    methodology = [
        (
            k,
            (
                f"Scope: {scope}. Processed archives: {', '.join(map(str, sorted(sources)))}. "
                "Other configured years remain NOT_PROCESSED. Pilot aggregates are not national estimates."
            )
            if k == "Scope"
            else (
                "Employer identifier gaps in processed years: "
                + (", ".join(gaps) or "none detected at whole-file level")
                + ". Whole-file gaps leave primary estimates missing. Validated-name results are separate sensitivities."
            )
            if k == "2017 gap"
            else v,
        )
        for k, v in methodology
    ]
    methodology = [
        ("Employer identifier gaps" if k == "2017 gap" else k, v) for k, v in methodology
    ]
    write_json(
        output / "analysis_manifest.json",
        {
            "classification_fingerprint": run["fingerprint"],
            "ads_hash": sha256(output / "ads_classified.parquet"),
            "config_hash": sha256(config_path),
            "analysis_modules": {
                name: sha256(Path(__file__).with_name(name))
                for name in ("aggregate.py", "export.py", "validate.py")
            },
            "geographic_scope": scope,
            "manual_validation": "pending"
            if not validation.manual_required.astype(str).str.strip().ne("").any()
            else "partially_or_fully_reviewed",
        },
    )
    tables = {
        "Municipality_Year": panel,
        "National_Trends": national,
        "Occupation_Trends": occupations,
        "Language_Categories": categories,
        "Validation": metrics(validation),
        "False_Negatives": false_negative_metrics(negatives),
        "AF_Comparison": af,
        "Coverage": panel[
            [
                "municipality_id",
                "municipality_name",
                "year",
                "occupation_group",
                "n_ads",
                "n_recruitment_spells",
                "coverage_flag",
                "year_status",
                "n_candidate_ads",
                "n_context_mixed_candidates",
                "n_context_uncertain_candidates",
                "n_unresolved_employer_candidates",
            ]
        ],
        "Employer_Matches": employer,
        "Methodology": pd.DataFrame(methodology, columns=["topic", "definition"]),
        "Audit": audit,
        "Spell_Trends": national_spells,
        "Sensitivity": pd.DataFrame(sensitivities),
    }
    table_dir = output / "tables"
    table_dir.mkdir(exist_ok=True)
    for name, frame in tables.items():
        frame.to_csv(table_dir / f"{name}.csv", index=False, encoding="utf-8-sig")
    if not discordant.empty:
        preserve_review(output / "af_discordant.csv", discordant)
        preserve_review(
            output / "af_discordant_sample.csv",
            discordant.sample(n=min(50, len(discordant)), random_state=config["seed"]),
        )
    write_json(
        output / "workbook_tables.json",
        {k: json.loads(v.to_json(orient="records", force_ascii=False)) for k, v in tables.items()},
    )
    write_json(
        output / "review_workbook_tables.json",
        {
            "Validation": json.loads(validation.to_json(orient="records", force_ascii=False)),
            "False_Negatives": json.loads(negatives.to_json(orient="records", force_ascii=False)),
            "Instructions": [
                {
                    "instruction": "Row height displays a preview only. Read the complete original text in the cell editor/formula bar or companion CSV before coding. Enter manual_required and manual_formal as true or false only when adjudicated. Leave uncertain cases blank with a note. Recruitment wording cannot establish municipal policy."
                }
            ],
        },
    )
    return tables, {
        "Validation": validation,
        "False_Negatives": negatives,
        "Instructions": pd.DataFrame([{"instruction": METHODOLOGY[16][1]}]),
    }


def export_xlsx(tables, path, *, preserve_existing=False):
    """Portable command-line exporter; refuse replacement of manual review workbooks."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    path = Path(path)
    if preserve_existing and path.exists():
        raise FileExistsError(
            f"Manual review workbook exists: {path}; preserve it and import labels first"
        )
    wb = Workbook()
    wb.remove(wb.active)
    for name, frame in tables.items():
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.append(list(frame.columns))
        for row in frame.itertuples(index=False, name=None):
            ws.append([None if pd.isna(v) else v for v in row])
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="243746")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 42
        ws.freeze_panes = "D2" if "municipality_id" in frame.columns else "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, column in enumerate(frame.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = (
                65
                if column
                in ("definition", "full_relevant_context", "matched_sentence", "instruction")
                else min(32, max(16, len(column) + 2))
            )
            for cells in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell = cells[0]
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Treat advertisement text as text, including '='.
                cell.font = Font(name="Arial", size=10, color="172B3A")
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=column in ("definition", "full_relevant_context", "instruction"),
                )
                if column.startswith("share_") or column.endswith("_share_required"):
                    cell.number_format = "0.0%"
                if column.startswith("manual_") or column == "notes":
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
            if column in ("manual_required", "manual_formal"):
                validation = DataValidation(type="list", formula1='"true,false"', allow_blank=True)
                ws.add_data_validation(validation)
                validation.add(
                    f"{get_column_letter(i)}2:{get_column_letter(i)}{max(2, ws.max_row)}"
                )
    wb.save(path)


def figures(tables, output):
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.ticker import PercentFormatter

    output = Path(output) / "figures"
    output.mkdir(exist_ok=True)
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 10,
            "axes.spines.top": False,
            "axes.spines.right": False,
        }
    )
    national, occupations, panel = (
        tables[k] for k in ("National_Trends", "Occupation_Trends", "Municipality_Year")
    )
    is_pilot = not (not national.empty and national.all_municipalities_in_scope.all())
    subtitle = (
        "Pilot municipalities only" if is_pilot else "Municipal Platsbanken recruitment"
    ) + " · unvalidated classifier estimates"
    gap_years = sorted(panel.loc[panel.coverage_flag.eq("EMPLOYER_ID_GAP"), "year"].unique())
    gap_note = (
        (", ".join(map(str, gap_years)) + " omitted: employer identifier gap. ")
        if gap_years
        else ""
    )

    def save(fig, name):
        fig.text(
            0.01,
            0.01,
            subtitle + "\n" + gap_note + "No formal-policy inference.",
            fontsize=8,
        )
        fig.tight_layout(rect=(0, 0.10, 1, 1))
        fig.savefig(output / f"{name}.png", dpi=160)
        fig.savefig(output / f"{name}.svg")
        plt.close(fig)

    plots = [
        (
            "01_required",
            "Ads requiring Swedish",
            ["share_required_ads", "standardised_share_required"],
        ),
        (
            "03_formal_qualitative",
            "Formal and qualitative requirements",
            [
                "share_formal_threshold",
                "share_strong_qualitative",
                "share_functional",
                "share_generic",
            ],
        ),
        (
            "04_formal_categories",
            "Specific formal requirements",
            [
                "share_sfi",
                "share_swedish_course",
                "share_gers_b1",
                "share_gers_b2",
                "share_language_test",
            ],
        ),
    ]
    for name, title, fields in plots:
        fig, ax = plt.subplots(figsize=(9, 5))
        for field in fields:
            if field in national:
                ax.plot(
                    national.year,
                    national[field],
                    "o-",
                    label=field.removeprefix("share_").replace("_", " "),
                )
        ax.set(title=title, xlabel="Publication year", ylabel="Share of eligible ads")
        ax.set_xticks(national.year.tolist())
        ax.yaxis.set_major_formatter(PercentFormatter(1))
        ymax = (
            max(
                0.01,
                max(
                    (
                        float(national[f].max())
                        for f in fields
                        if f in national and not national.empty
                    ),
                    default=0,
                )
                * 1.25,
            )
            if name == "04_formal_categories"
            else 1
        )
        ax.set_ylim(0, ymax)
        ax.legend(fontsize=8)
        save(fig, name)
    fig, ax = plt.subplots(figsize=(9, 5))
    if not occupations.empty:
        for code, group in occupations.groupby("occupation_group"):
            ax.plot(group.year, group.share_required_ads, "o-", label=f"SSYK {code}")
    ax.set(
        title="Requirements by occupation",
        xlabel="Publication year",
        ylabel="Share of eligible ads",
        ylim=(0, 1),
    )
    ax.set_xticks(sorted(occupations.year.unique()))
    ax.yaxis.set_major_formatter(PercentFormatter(1))
    ax.legend()
    save(fig, "02_occupations")
    fig, ax = plt.subplots(figsize=(9, 5))
    observed = panel[(panel.occupation_group == "overall") & panel.complete_year_trend_eligible]
    for year, group in observed.groupby("year"):
        ax.scatter(
            [year] * group.share_required_ads.notna().sum(),
            group.share_required_ads.dropna(),
            label=str(year),
            alpha=0.7,
        )
    ax.set(
        title="Municipality variation in observed requirement shares",
        xlabel="Publication year",
        ylabel="Municipality share",
        ylim=(0, 1),
    )
    ax.yaxis.set_major_formatter(PercentFormatter(1))
    ax.set_xticks(sorted(observed.year.unique()))
    save(fig, "05_municipality_distribution")
    fig, ax = plt.subplots(figsize=(9, 5))
    counts = observed.assign(good=observed.coverage_flag.eq("GOOD")).groupby("year").good.sum()
    ax.bar(counts.index.astype(str), counts.values, color="#315E76")
    ax.set(
        title="Municipalities in scope with at least 20 eligible ads",
        xlabel="Publication year",
        ylabel="Number of municipalities",
    )
    save(fig, "06_coverage")
