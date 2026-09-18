from pathlib import Path

from openpyxl import load_workbook

from municipal_research.cli import main
from municipal_research.config import load_config
from municipal_research.demo import run_demo
from municipal_research.export import export_workbook
from municipal_research.pipeline import summarize
from municipal_research.storage import read_json, write_json


def test_full_html_pdf_pipeline_and_offline_replay(config, tmp_path):
    run, cache = tmp_path / "run", tmp_path / "cache"
    data = run_demo(config, run, cache)
    assert [s["category"] for s in data["summary"]] == ["Yes", "Pilot"]
    assert not data["errors"]
    assert {d["media_type"] for d in data["documents"]} == {"html", "pdf"}
    docs = {d["id"]: d for d in data["documents"]}
    for result in data["classifications"]:
        text = (run / docs[result["document_id"]]["text_path"]).read_bytes().decode("utf-8")
        for quote in result["quotes"]:
            for loc in quote["locations"]:
                assert text[loc["start"] : loc["end"]] == loc["matched_text"]
    assert main(["verify-run", str(run)]) == 0
    replay = run_demo(config, tmp_path / "replay", cache, offline=True)
    assert replay["classifications"] == data["classifications"]
    assert read_json(tmp_path / "replay/manifest.json")["api_calls"] == 0
    workbook = load_workbook(run / "results.xlsx")
    assert len(workbook.sheetnames) == 10
    assert workbook["Summary"].freeze_panes == "C2"
    assert workbook["Summary"].max_row == 3
    for sheet in workbook:
        assert sheet.max_column >= 2
        assert all(cell.data_type not in {"f", "e"} for row in sheet for cell in row)
    workbook.close()
    first_text = run / data["documents"][0]["text_path"]
    first_text.write_text("tampered", encoding="utf-8")
    assert main(["verify-run", str(run)]) == 1


def test_excel_injection_is_literal(config, tmp_path):
    run = tmp_path / "run"
    data = run_demo(config, run, tmp_path / "cache")
    data["summary"][0]["municipality"] = '=HYPERLINK("https://evil.example","click")'
    write_json(run / "results.json", data)
    workbook_path = export_workbook(run, tmp_path / "export.xlsx")
    book = load_workbook(workbook_path)
    assert book["Summary"]["B2"].data_type == "s"
    assert book["Summary"]["B2"].value.startswith("=HYPERLINK")
    book.close()


def test_summary_preserves_conflicts_and_no_evidence(config, municipality, yes_decision):
    base = {"decision": yes_decision.model_dump(), "needs_review": False, "quotes": []}
    other = {**base, "decision": {**base["decision"], "category": "Pilot"}}
    result = summarize(municipality, [], [base, other], [], config)
    assert result["category"] == "Review" and result["needs_review"]
    empty = summarize(municipality, [], [], ["download_failed"], config)
    assert empty["category"] == "NoEvidence" and empty["needs_review"]


def test_generic_configuration():
    config = load_config(Path(__file__).parents[1] / "examples/generic.yaml")
    assert config.research.cutoff is None
    assert set(config.research.labels) == {"Target", "Unknown"}
